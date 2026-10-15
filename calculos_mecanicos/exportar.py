"""
Gráficas y exportaciones: Excel (memoria de cálculo), PDF (plantillado) y DXF (1:1).

Mejoras frente al original:
  * Las catenarias se calculan UNA vez y se reutilizan (antes se recalculaban 4 veces).
  * Una sola función dibuja el perfil: la usan la pantalla y el PDF (antes había
    dos copias del mismo código).
  * Se dibuja terreno, línea de distancia mínima de seguridad y postes.
  * El usuario elige la carpeta de salida; los archivos no se sobrescriben en
    silencio dentro de la carpeta del programa.
  * DXF con ezdxf si está instalado; si no, escritor DXF R12 propio (sin dependencias).
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

from .cambio_estado import tabla_tendido
from .plantillado import ResultadoPlantillado, tabla_secciones


# ===========================================================================
# GRÁFICA DEL PERFIL
# ===========================================================================
def abscisa_k(m: float) -> str:
    """Formato de abscisa vial: 1780.5 -> 'K1+780.50'."""
    return f"K{int(m // 1000)}+{m % 1000:06.2f}"


def paginas(res: ResultadoPlantillado) -> list[tuple[int, int]]:
    n = len(res.catenarias)
    paso = max(1, res.proyecto.salida.vanos_por_pagina)
    return [(i, min(i + paso, n)) for i in range(0, n, paso)]


# Escalas normalizadas de plano (1:N)
_ESCALAS = [100, 125, 200, 250, 500, 750, 1000, 1250, 1500, 2000, 2500, 5000, 7500, 10000, 20000, 25000, 50000]


def _rango_y(res: ResultadoPlantillado, i: int, j: int) -> tuple[float, float]:
    """Cotas mínima y máxima a dibujar para los vanos [i, j)."""
    ys = []
    for c, cf in zip(res.catenarias[i:j], res.catenarias_frias[i:j]):
        ys += [float(c.y1), float(c.y2), float(c.y(c.x_flecha_maxima)), float(cf.y(cf.x_flecha_maxima))]
    ap = res.apoyos
    ys += list(ap["Cota Y (m)"].iloc[i:j + 1]) + list(ap["Cota amarre (m)"].iloc[i:j + 1])
    if res.datos.terreno is not None or res.datos.y_es_terreno:
        xt = np.linspace(ap["X (m)"].iloc[i], ap["X (m)"].iloc[j], 30 * (j - i) + 1)
        yt = res.terreno(xt)
        ys += [float(yt.min()), float(yt.max() + res.proyecto.criterios.distancia_minima_terreno_m)]
    return min(ys), max(ys)


def paginas_escala(res: ResultadoPlantillado, ancho_in: float, alto_in: float, ex: float):
    """Tramos del perfil a escala fija X 1:N · Y 1:(N/ex) (misma N en todas las hojas).

    Devuelve (tramos, N). N es la escala normalizada más pequeña con la que caben ~vanos_por_pagina
    vanos típicos; si un tramo no cabe a lo ancho o a lo alto, se parte.
    """
    ap = res.apoyos["X (m)"].to_numpy()
    vanos = np.sort(np.diff(ap))
    mediana = float(vanos[len(vanos) // 2]) if len(vanos) else 100.0
    vpp = max(1, res.proyecto.salida.vanos_por_pagina)
    n_req = mediana * vpp / (ancho_in * 0.0254)        # 1:N para que quepan vpp vanos típicos
    ref = min(vpp, len(vanos))                          # …y que su desnivel típico quepa a lo alto
    desn = sorted(_rango_y(res, i, i + ref)[1] - _rango_y(res, i, i + ref)[0] + 6
                  for i in range(0, len(vanos) - ref + 1)) or [20.0]
    n_req = max(n_req, desn[len(desn) // 2] * ex / (alto_in * 0.0254))
    N = next((e for e in _ESCALAS if e >= n_req), _ESCALAS[-1])
    m_por_in = N * 0.0254                               # metros reales por pulgada de papel
    tramos, i, n = [], 0, len(vanos)
    while i < n:
        j = i + 1
        while j < n:
            ancho = (ap[j + 1] - ap[i]) / m_por_in
            lo, hi = _rango_y(res, i, j + 1)
            alto = (hi - lo + 6) * ex / m_por_in
            if ancho > ancho_in or alto > alto_in:
                break
            j += 1
        tramos.append((i, j))
        i = j
    return tramos, N


def dibujar_perfil(ax, res: ResultadoPlantillado, inicio: int, fin: int, escala=None) -> None:
    """Dibuja los vanos [inicio, fin) en el eje matplotlib `ax`."""
    ap = res.apoyos
    x0 = ap["X (m)"].iloc[0]
    dmin = res.proyecto.criterios.distancia_minima_terreno_m
    npts = res.proyecto.salida.puntos_por_vano
    ys_all = []

    xa, xb = ap["X (m)"].iloc[inicio], ap["X (m)"].iloc[fin]
    tiene_terreno = res.datos.terreno is not None or res.datos.y_es_terreno
    if tiene_terreno:
        xt = np.linspace(xa, xb, 50 * (fin - inicio) + 1)
        yt = res.terreno(xt)
        ax.fill_between(xt, yt, yt.min() - 1000, color="#d8c7a8", alpha=0.6, lw=0)
        ax.plot(xt, yt, color="#7a5c2e", lw=1.2, label="Terreno")
        ax.plot(xt, yt + dmin, color="#c0392b", lw=1, ls="--",
                label=f"Distancia mínima {dmin:.2f} m")
        ys_all += [yt.min(), (yt + dmin).max()]

    for i in range(inicio, fin):
        c = res.catenarias[i]
        x, y = c.muestrear(npts)
        ax.plot(x, y, color="#1f5aa6", lw=2,
                label="Conductor (plantilla caliente)" if i == inicio else None)
        if res.proyecto.modo_flecha == "cambio_estado":
            cf = res.catenarias_frias[i]
            xf, yf = cf.muestrear(npts)
            ax.plot(xf, yf, color="#5dade2", lw=0.9, ls=":",
                    label="Conductor (plantilla fría)" if i == inicio else None)
        ys_all += [y.min(), y.max()]
        fila = res.vanos.iloc[i]
        ax.annotate(f"a={c.vano:.1f} m\nf={c.flecha:.2f} m",
                    ((c.x1 + c.x2) / 2, c.y((c.x1 + c.x2) / 2)),
                    xytext=(0, -26), textcoords="offset points", ha="center",
                    fontsize=7, color="#1f5aa6")
        if fila.get("Cumple distancia") == "NO":
            ax.plot(fila["X dist. mín."], c.y(fila["X dist. mín."]), "x", color="red", ms=10, mew=2)

    for j in range(inicio, fin + 1):
        r = ap.iloc[j]
        xs, ya, yt = r["X (m)"], r["Cota amarre (m)"], r["Cota Y (m)"]
        if tiene_terreno and res.datos.y_es_terreno:
            ax.plot([xs, xs], [yt, ya], color="#333", lw=2.5, solid_capstyle="butt")
        mark = "s" if r["Tipo"] == "Retención" else "o"
        ax.plot(xs, ya, mark, color="#c0392b", ms=6, zorder=5)
        txt = f"{r['Apoyo']}" + (" ▲" if r["Arrancamiento"] == "SÍ" else "")
        ax.annotate(txt, (xs, ya), xytext=(0, 8), textcoords="offset points",
                    ha="center", fontsize=8, fontweight="bold")
        ys_all.append(ya)

    ax.set_xticks(ap["X (m)"].iloc[inicio:fin + 1])
    ax.set_xticklabels([abscisa_k(v - x0) for v in ap["X (m)"].iloc[inicio:fin + 1]], fontsize=7, rotation=0)
    lo, hi = min(ys_all), max(ys_all)
    marg = max(3.0, 0.15 * (hi - lo))
    ax.set_ylim(lo - marg, hi + marg)
    ax.set_xlim(xa - 0.01 * (xb - xa), xb + 0.01 * (xb - xa))
    if escala:  # (N, ex, ancho_in, alto_in): misma escala en todas las hojas
        N, ex, ancho_in, alto_in = escala
        m_por_in = N * 0.0254
        lo_d, hi_d = _rango_y(res, inicio, fin)
        ax.set_xlim(xa, xa + ancho_in * m_por_in)
        alto_m = alto_in * m_por_in / ex
        yc = (lo_d + hi_d) / 2
        ax.set_ylim(yc - alto_m / 2, yc + alto_m / 2)
    paso_y = float(getattr(res.proyecto.salida, "intervalo_cotas_m", 0) or 0)
    if paso_y > 0:  # cotas cada `paso_y` m (5 m por defecto); rótulos espaciados si quedan muy juntos
        from matplotlib.ticker import MultipleLocator
        lo_y, hi_y = ax.get_ylim()
        n = max(1, int(-(-((hi_y - lo_y) / paso_y) // 40)))   # máx. ~40 rótulos
        ax.yaxis.set_major_locator(MultipleLocator(paso_y * n))
        ax.yaxis.set_minor_locator(MultipleLocator(paso_y))
        ax.grid(True, which="minor", axis="y", alpha=0.15)
    ax.grid(True, alpha=0.35)
    ax.set_xlabel("Abscisa (K + m)")
    ax.set_ylabel("Cota (m)")
    ax.set_title(f"Plantillado – vanos {inicio + 1} a {fin}"
                 + (f"   ·   Escala horizontal 1:{escala[0]:g} · vertical 1:{escala[0] / escala[1]:g}"
                    f" (Y {escala[1]:g}:1)" if escala else ""), fontsize=11)
    ax.legend(loc="upper right", fontsize=7, framealpha=0.85)


# ===========================================================================
# EXCEL
# ===========================================================================
def _tabla_parametros(res: ResultadoPlantillado) -> pd.DataFrame:
    p = res.proyecto
    filas = [("Proyecto", p.nombre), ("Fecha de cálculo", datetime.now().strftime("%Y-%m-%d %H:%M"))]
    for k, v in vars(p.conductor).items():
        filas.append((f"Conductor · {k}", v))
    for k, h in p.hipotesis.items():
        filas.append((f"Hipótesis {k}",
                      f"{h.nombre}: {h.temperatura_c} °C, viento {h.viento_kmh} km/h "
                      f"(q={h.presion_viento_kgf_m2():.1f} kgf/m²), hielo {h.hielo_kgf_m} kgf/m, "
                      f"w={h.peso_aparente(p.conductor):.4f} kgf/m, "
                      f"límite={'—' if not h.limite_fraccion_rotura else f'{100*h.limite_fraccion_rotura:.0f} % rotura'}"))
    for k, v in vars(p.criterios).items():
        filas.append((f"Criterio · {k}", v))
    filas += [("Modo de flecha", p.modo_flecha),
              ("Hipótesis de plantillado", p.hipotesis_plantillado),
              ("Hipótesis de arrancamiento", p.hipotesis_arrancamiento)]
    return pd.DataFrame(filas, columns=["Parámetro", "Valor"])


def exportar_excel(res: ResultadoPlantillado, ruta: str | Path) -> Path:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ruta = Path(ruta)
    p = res.proyecto
    hojas = {
        "Resumen": pd.DataFrame(list(res.resumen.items()), columns=["Concepto", "Valor"]),
        "Vanos": res.vanos,
        "Apoyos": res.apoyos,
        "Secciones": tabla_secciones(res),
    }
    if getattr(res, "postes", None) is not None:
        hojas["Postes y templetes"] = res.postes
    if p.modo_flecha == "cambio_estado":
        hojas["Tabla de tendido"] = tabla_tendido(p.conductor, p.hipotesis,
                                                  list(range(20, 405, 5)),
                                                  p.criterios.fluencia_delta_t_c)
    hojas["Advertencias"] = pd.DataFrame({"Advertencia": res.advertencias or ["Sin advertencias"]})
    hojas["Parámetros"] = _tabla_parametros(res)

    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        for nombre, df in hojas.items():
            df.to_excel(xw, sheet_name=nombre, index=False)
            ws = xw.sheets[nombre]
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F5AA6")
                c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            ws.freeze_panes = "A2"
            for j, col in enumerate(df.columns, start=1):
                largo = max([len(str(col))] + [len(str(v)) for v in df[col].head(300)])
                ws.column_dimensions[get_column_letter(j)].width = min(max(10, largo + 2), 90)
            rojo = PatternFill("solid", fgColor="F8D7DA")
            enc = {c.value: c.column for c in ws[1]}
            alertas = [(enc[k], v) for k, v in (("Cumple distancia", "NO"), ("Arrancamiento", "SÍ"),
                                                         ("Estado", "REQUIERE TEMPLETE"), ("Estado", "TEMPLETE INSUFICIENTE"))
                       if k in enc]
            for fila in ws.iter_rows(min_row=2):
                if any(fila[col - 1].value == v for col, v in alertas):
                    for c in fila:
                        c.fill = rojo
    return ruta


# ===========================================================================
# PDF
# ===========================================================================
def exportar_pdf(res: ResultadoPlantillado, ruta: str | Path) -> Path:
    import matplotlib
    matplotlib.use("Agg", force=False)
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    ruta = Path(ruta)
    with PdfPages(ruta) as pdf:
        # --- portada / resumen
        fig = plt.figure(figsize=(16.54, 11.69))  # A3 horizontal
        fig.text(0.05, 0.93, res.proyecto.nombre, fontsize=20, fontweight="bold")
        fig.text(0.05, 0.89, "Memoria de plantillado y cálculo mecánico de conductores",
                 fontsize=13)
        ax = fig.add_axes([0.05, 0.38, 0.42, 0.48]); ax.axis("off")
        t = ax.table(cellText=[[k, str(v)] for k, v in res.resumen.items() if k != "Archivo"],
                     colLabels=["Concepto", "Valor"], loc="upper left", cellLoc="left",
                     colWidths=[0.62, 0.38])
        t.auto_set_font_size(False); t.set_fontsize(9); t.scale(1, 1.4)
        ax2 = fig.add_axes([0.52, 0.38, 0.45, 0.48]); ax2.axis("off")
        sec = tabla_secciones(res)
        cols = [c for c in sec.columns if not c.startswith("%")][:9]
        t2 = ax2.table(cellText=sec[cols].astype(str).values.tolist(), colLabels=cols,
                       loc="upper left", cellLoc="center")
        t2.auto_set_font_size(False); t2.set_fontsize(7.5); t2.scale(1, 1.4)
        adv = res.advertencias[:22]
        fig.text(0.05, 0.33, "Advertencias", fontsize=12, fontweight="bold")
        fig.text(0.05, 0.31, "\n".join(f"• {a[:190]}" for a in adv) or "Sin advertencias",
                 fontsize=7.5, va="top")
        pdf.savefig(fig); plt.close(fig)

        # --- perfiles
        ex = float(getattr(res.proyecto.salida, "exageracion_vertical_perfil", 0) or 0)
        caja = [0.06, 0.08, 0.92, 0.85]                     # área del perfil en la hoja A3
        ancho_in, alto_in = 16.54 * caja[2], 11.69 * caja[3]
        if ex > 0:
            tramos, N = paginas_escala(res, ancho_in, alto_in, ex)
        else:
            tramos, N = paginas(res), None
        for ini, fin in tramos:
            fig = plt.figure(figsize=(16.54, 11.69))
            if N:
                ax = fig.add_axes(caja)
                dibujar_perfil(ax, res, ini, fin, (N, ex, ancho_in, alto_in))
            else:
                ax = fig.add_subplot(111)
                dibujar_perfil(ax, res, ini, fin)
                fig.tight_layout()
            pdf.savefig(fig); plt.close(fig)

        # --- tabla de vanos
        cols = ["Vano", "Apoyo inicial", "Apoyo final", "Sección", "Vano a (m)", "Desnivel (m)",
                "Flecha máx. (m)", "Longitud conductor (m)", "T apoyo ini (kgf)",
                "T apoyo fin (kgf)", "Separación mín. fases (m)"]
        cols += [c for c in ("Dist. mín. terreno (m)", "Cumple distancia") if c in res.vanos]
        for i in range(0, len(res.vanos), 30):
            fig = plt.figure(figsize=(16.54, 11.69))
            ax = fig.add_axes([0.03, 0.03, 0.94, 0.9]); ax.axis("off")
            ax.set_title("Resultados por vano", fontsize=13, loc="left")
            tb = ax.table(cellText=res.vanos[cols].iloc[i:i + 30].astype(str).values.tolist(),
                          colLabels=cols, loc="upper center", cellLoc="center")
            tb.auto_set_font_size(False); tb.set_fontsize(7.5); tb.scale(1, 1.5)
            pdf.savefig(fig); plt.close(fig)
    return ruta


# ===========================================================================
# DXF
# ===========================================================================
_CAPAS = {"TERRENO": 30, "DIST_SEGURIDAD": 1, "CONDUCTOR": 5, "CONDUCTOR_FRIO": 4,
          "POSTES": 7, "APOYOS": 1, "TEXTOS_APOYOS": 2, "ABSCISAS": 3, "FLECHAS": 6}


class _DXFR12:
    """Escritor DXF R12 mínimo (sin dependencias). Abre en AutoCAD/BricsCAD/QCAD."""

    def __init__(self):
        self.ent: list[str] = []

    def linea(self, p1, p2, capa):
        self.ent += ["0", "LINE", "8", capa, "10", f"{p1[0]:.4f}", "20", f"{p1[1]:.4f}", "30", "0",
                     "11", f"{p2[0]:.4f}", "21", f"{p2[1]:.4f}", "31", "0"]

    def polilinea(self, pts, capa):
        self.ent += ["0", "POLYLINE", "8", capa, "66", "1", "10", "0", "20", "0", "30", "0", "70", "0"]
        for x, y in pts:
            self.ent += ["0", "VERTEX", "8", capa, "10", f"{x:.4f}", "20", f"{y:.4f}", "30", "0"]
        self.ent += ["0", "SEQEND", "8", capa]

    def circulo(self, c, r, capa):
        self.ent += ["0", "CIRCLE", "8", capa, "10", f"{c[0]:.4f}", "20", f"{c[1]:.4f}", "30", "0",
                     "40", f"{r:.4f}"]

    def texto(self, txt, p, h, capa, centrado=True):
        e = ["0", "TEXT", "8", capa, "10", f"{p[0]:.4f}", "20", f"{p[1]:.4f}", "30", "0",
             "40", f"{h:.4f}", "1", txt]
        if centrado:
            e += ["72", "1", "11", f"{p[0]:.4f}", "21", f"{p[1]:.4f}", "31", "0"]
        self.ent += e

    def guardar(self, ruta):
        capas = ["0", "TABLE", "2", "LAYER", "70", str(len(_CAPAS))]
        for n, col in _CAPAS.items():
            capas += ["0", "LAYER", "2", n, "70", "0", "62", str(col), "6", "CONTINUOUS"]
        capas += ["0", "ENDTAB"]
        txt = (["0", "SECTION", "2", "HEADER", "9", "$ACADVER", "1", "AC1009", "0", "ENDSEC",
                "0", "SECTION", "2", "TABLES"] + capas + ["0", "ENDSEC",
                "0", "SECTION", "2", "ENTITIES"] + self.ent + ["0", "ENDSEC", "0", "EOF"])
        Path(ruta).write_text("\n".join(txt) + "\n", encoding="cp1252", errors="replace")


class _DXFEzdxf:
    def __init__(self):
        import ezdxf
        self.doc = ezdxf.new("R2018")
        self.msp = self.doc.modelspace()
        for n, col in _CAPAS.items():
            if n not in self.doc.layers:
                self.doc.layers.add(n, color=col)

    def linea(self, p1, p2, capa):
        self.msp.add_line(p1, p2, dxfattribs={"layer": capa})

    def polilinea(self, pts, capa):
        self.msp.add_lwpolyline(pts, dxfattribs={"layer": capa})

    def circulo(self, c, r, capa):
        self.msp.add_circle(c, r, dxfattribs={"layer": capa})

    def texto(self, txt, p, h, capa, centrado=True):
        from ezdxf.enums import TextEntityAlignment
        t = self.msp.add_text(txt, height=h, dxfattribs={"layer": capa})
        t.set_placement(p, align=TextEntityAlignment.BOTTOM_CENTER if centrado
                        else TextEntityAlignment.LEFT)

    def guardar(self, ruta):
        self.doc.saveas(ruta)


def exportar_dxf(res: ResultadoPlantillado, ruta: str | Path, forzar_r12: bool = False) -> Path:
    ruta = Path(ruta)
    try:
        if forzar_r12:
            raise ImportError
        d = _DXFEzdxf()
    except ImportError:
        d = _DXFR12()

    s = res.proyecto.salida
    ev = s.exageracion_vertical_dxf
    h = s.altura_texto_dxf
    ap = res.apoyos
    x0 = ap["X (m)"].iloc[0]
    yref = float(ap["Cota amarre (m)"].min())
    Y = lambda y: yref + (y - yref) * ev       # exageración vertical (1 = escala 1:1)
    P = lambda x, y: (float(x), float(Y(y)))

    # terreno y línea de seguridad
    if res.datos.terreno is not None or res.datos.y_es_terreno:
        xt = np.linspace(ap["X (m)"].iloc[0], ap["X (m)"].iloc[-1], 20 * len(ap) + 1)
        if res.datos.terreno is not None:
            xt = np.union1d(xt, res.datos.terreno["x"].to_numpy())
        yt = res.terreno(xt)
        d.polilinea([P(a, b) for a, b in zip(xt, yt)], "TERRENO")
        dmin = res.proyecto.criterios.distancia_minima_terreno_m
        d.polilinea([P(a, b + dmin) for a, b in zip(xt, yt)], "DIST_SEGURIDAD")

    # conductores
    for c, cf in zip(res.catenarias, res.catenarias_frias):
        x, y = c.muestrear(s.puntos_por_vano)
        d.polilinea([P(a, b) for a, b in zip(x, y)], "CONDUCTOR")
        if res.proyecto.modo_flecha == "cambio_estado":
            x, y = cf.muestrear(s.puntos_por_vano)
            d.polilinea([P(a, b) for a, b in zip(x, y)], "CONDUCTOR_FRIO")
        xm = c.x_flecha_maxima
        d.linea(P(xm, c.cuerda(xm)), P(xm, c.y(xm)), "FLECHAS")
        d.texto(f"a={c.vano:.2f} f={c.flecha:.2f}", P((c.x1 + c.x2) / 2, c.y(xm) - 2.5 * h / ev),
                h * 0.7, "FLECHAS")

    # apoyos, postes, textos, abscisas
    for _, r in ap.iterrows():
        xs, ya, yt = r["X (m)"], r["Cota amarre (m)"], r["Cota Y (m)"]
        if res.datos.y_es_terreno:
            d.linea(P(xs, yt), P(xs, ya), "POSTES")
        d.circulo(P(xs, ya), 0.3, "APOYOS")
        d.texto(str(r["Apoyo"]), (xs, Y(ya) + 1.5 * h), h, "TEXTOS_APOYOS")
        ab = xs - x0
        base = Y(min(yt, ya)) - 2.5 * h
        d.texto(abscisa_k(ab), (xs, base), h * 0.8, "ABSCISAS")
    d.guardar(ruta)
    return ruta


def exportar_todo(res: ResultadoPlantillado, carpeta: str | Path, prefijo: str = "PLANTILLADO"):
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    return {
        "xlsx": exportar_excel(res, carpeta / f"{prefijo}_RESULTADOS.xlsx"),
        "pdf": exportar_pdf(res, carpeta / f"{prefijo}_PERFIL.pdf"),
        "dxf": exportar_dxf(res, carpeta / f"{prefijo}_PERFIL.dxf"),
    }
